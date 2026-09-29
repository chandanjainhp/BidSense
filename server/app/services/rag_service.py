"""RAG (Retrieval-Augmented Generation) service.

Pipeline:
    upload -> extract text (pdf/docx/txt/md) -> chunk with overlap -> embed chunks
    -> store in Postgres (pgvector, HNSW index) -> hybrid retrieval (semantic + keyword)
    -> context-augmented LLM answer with cited sources.

Works in two modes:
    * online  -- any OpenAI-compatible /embeddings + /chat/completions endpoint
    * offline -- deterministic hashed bag-of-words embeddings (no API key needed),
      so search still works for local/dev setups; quality is lower but functional.
"""

import hashlib
import math
import re
import uuid

import httpx
from sqlalchemy import select, delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.models.document import Document, DocumentChunk, DocumentStatus

# ---------------------------------------------------------------------------
# Tunables (mirrored in docs/RAG_SYSTEM.md)
# ---------------------------------------------------------------------------
CHUNK_SIZE = 1000        # target characters per chunk
CHUNK_OVERLAP = 200      # characters of overlap between consecutive chunks
MIN_CHUNK_LEN = 40       # chunks shorter than this are dropped
TOP_K_DEFAULT = 5
SEMANTIC_CANDIDATES = 25  # candidates pulled from vector search before blending
KEYWORD_CANDIDATES = 25  # candidates pulled from keyword search before blending
HYBRID_ALPHA = 0.7       # weight of semantic score in the blend (1.0 = pure semantic)
MAX_CONTEXT_CHARS = 6000  # cap on the injected context block
EMBED_DIM = 1536         # must match DocumentChunk.embedding column dimension

EMBEDDING_BATCH_SIZE = 32

SYSTEM_PROMPT = (
    "You are BidSense AI, a procurement and RFP management assistant. "
    "Answer the user's question using primarily the KNOWLEDGE BASE CONTEXT provided. "
    "Cite the documents you used inline in the format [source: filename] "
    "whenever a statement comes from the context. "
    "If the context does not contain the answer, say so plainly instead of guessing."
)


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------
def extract_text(filename: str, raw: bytes) -> str:
    """Extract plain text from an uploaded file (pdf, docx, txt, md)."""
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "pdf":
        return _extract_pdf(raw)
    if suffix == "docx":
        return _extract_docx(raw)
    if suffix in ("txt", "md"):
        return raw.decode("utf-8", errors="replace")
    # Best-effort fallback: treat as plain text
    return raw.decode("utf-8", errors="replace")


def _extract_pdf(raw: bytes) -> str:
    from io import BytesIO
    from pdfminer.high_level import extract_text as pdf_extract_text

    return pdf_extract_text(BytesIO(raw))


def _extract_docx(raw: bytes) -> str:
    from io import BytesIO
    import docx  # python-docx

    doc = docx.Document(BytesIO(raw))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------
def chunk_text(text: str) -> list[str]:
    """Split text into overlapping chunks on paragraph boundaries where possible.

    Paragraph-aware splitting keeps semantic units together; long paragraphs are
    split on sentence boundaries; sentences longer than CHUNK_SIZE are hard-split.
    """
    text = text.strip()
    if not text:
        return []

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]

    units: list[str] = []  # paragraph or sub-paragraph pieces <= CHUNK_SIZE
    for para in paragraphs:
        if len(para) <= CHUNK_SIZE:
            units.append(para)
            continue
        # Split long paragraph into sentences
        sentences = re.split(r"(?<=[.!?])\s+", para)
        current = ""
        for sentence in sentences:
            if len(current) + len(sentence) + 1 <= CHUNK_SIZE:
                current = f"{current} {sentence}".strip()
            else:
                if current:
                    units.append(current)
                # Hard-split pathological sentences
                while len(sentence) > CHUNK_SIZE:
                    units.append(sentence[:CHUNK_SIZE])
                    sentence = sentence[CHUNK_SIZE:]
                current = sentence
        if current:
            units.append(current)

    # Assemble units into chunks of ~CHUNK_SIZE with CHUNK_OVERLAP carry-over
    chunks: list[str] = []
    current = ""
    for unit in units:
        candidate = f"{current}\n{unit}".strip() if current else unit
        if len(candidate) > CHUNK_SIZE and current:
            chunks.append(current)
            # carry overlap tail from the previous chunk
            overlap_tail = current[-CHUNK_OVERLAP:]
            current = f"{overlap_tail}\n{unit}".strip()
            if len(current) > CHUNK_SIZE:
                chunks.append(current)
                current = ""
        else:
            current = candidate
    if current:
        chunks.append(current)

    return [c for c in chunks if len(c) >= MIN_CHUNK_LEN]


# ---------------------------------------------------------------------------
# Embeddings
# ---------------------------------------------------------------------------
def _tokenize(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", text.lower())


def _offline_embed(text: str) -> list[float]:
    """Deterministic hashed bag-of-words embedding (offline fallback).

    No API needed: each token is hashed into EMBED_DIM buckets; sublinear tf
    weighting + L2 normalization gives a usable similarity signal. Quality is
    below a real embedding model but search remains functional offline.
    """
    vec = [0.0] * EMBED_DIM
    tokens = _tokenize(text)
    if not tokens:
        return vec
    counts: dict[str, int] = {}
    for tok in tokens:
        counts[tok] = counts.get(tok, 0) + 1
    for tok, count in counts.items():
        bucket = int(hashlib.md5(tok.encode("utf-8")).hexdigest(), 16) % EMBED_DIM
        vec[bucket] += 1.0 + math.log(count)
    norm = math.sqrt(sum(v * v for v in vec))
    if norm > 0:
        vec = [v / norm for v in vec]
    return vec


class RagService:
    """Ingestion + retrieval + context building for the knowledge base."""

    def __init__(self):
        self.api_key = settings.AI_API_KEY
        self.base_url = getattr(settings, "AI_BASE_URL", "https://api.openai.com/v1")
        self.model = settings.AI_MODEL
        # Any OpenAI-compatible embedding model; used only when AI_API_KEY is set
        self.embedding_model = getattr(settings, "AI_EMBEDDING_MODEL", "text-embedding-3-small")
        self.online = bool(self.api_key)

    # -- embedding helpers ---------------------------------------------------
    @property
    def embed_dim(self) -> int:
        return EMBED_DIM

    async def embed_texts(self, texts: list[str]) -> list[list[float]]:
        """Embed a batch of texts. Uses the API when configured, else offline hashing."""
        if not texts:
            return []
        if self.online:
            try:
                async with httpx.AsyncClient() as client:
                    out: list[list[float]] = []
                    for i in range(0, len(texts), EMBEDDING_BATCH_SIZE):
                        batch = texts[i : i + EMBEDDING_BATCH_SIZE]
                        response = await client.post(
                            f"{self.base_url}/embeddings",
                            headers={
                                "Authorization": f"Bearer {self.api_key}",
                                "Content-Type": "application/json",
                            },
                            json={"model": self.embedding_model, "input": batch},
                            timeout=60.0,
                        )
                        response.raise_for_status()
                        data = response.json()["data"]
                        # API may return batches out of order; restore order by index
                        data.sort(key=lambda d: d["index"])
                        out.extend([d["embedding"] for d in data])
                    return out
            except Exception:
                pass  # fall through to offline embedding
        return [_offline_embed(t) for t in texts]

    # -- ingestion -----------------------------------------------------------
    async def ingest_document(
        self,
        db: AsyncSession,
        *,
        user_id: uuid.UUID,
        filename: str,
        raw: bytes,
    ) -> Document:
        """Full ingestion pipeline for one uploaded file. Commits via caller's session."""
        size_bytes = len(raw)
        doc = Document(
            user_id=user_id,
            filename=filename,
            file_type=filename.rsplit(".", 1)[-1].lower() if "." in filename else "raw",
            size_bytes=size_bytes,
            status=DocumentStatus.PROCESSING,
        )
        db.add(doc)
        await db.flush()  # assign doc.id

        try:
            text = extract_text(filename, raw)
            if not text or not text.strip():
                raise ValueError("No extractable text found in file")

            chunks = chunk_text(text)
            if not chunks:
                raise ValueError("Document produced no usable chunks")

            embeddings = await self.embed_texts(chunks)
            for idx, (chunk, emb) in enumerate(zip(chunks, embeddings)):
                db.add(
                    DocumentChunk(
                        document_id=doc.id,
                        chunk_index=idx,
                        content=chunk,
                        embedding=emb,
                    )
                )
            doc.chunk_count = len(chunks)
            doc.status = DocumentStatus.READY
            await db.flush()
        except Exception as exc:  # noqa: BLE001 - record failure, keep row for debugging
            doc.status = DocumentStatus.FAILED
            doc.error_message = str(exc)[:500]
            await db.flush()

        return doc

    async def delete_document(self, db: AsyncSession, doc: Document) -> None:
        await db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc.id))
        await db.delete(doc)

    # -- retrieval -----------------------------------------------------------
    async def retrieve(
        self,
        db: AsyncSession,
        *,
        user_id: uuid.UUID,
        query: str,
        top_k: int = TOP_K_DEFAULT,
        mode: str = "hybrid",
    ) -> list[dict]:
        """Retrieve the most relevant chunks for a query.

        Returns dicts: {document_id, filename, chunk_index, content, score}.
        Scores are normalized to roughly [0, 1] with higher = better.
        """
        if not query.strip():
            return []

        parts: dict[tuple, dict] = {}

        # 1) Semantic (pgvector cosine distance -> similarity)
        if mode in ("hybrid", "semantic"):
            query_emb = (await self.embed_texts([query]))[0]
            stmt = (
                select(
                    DocumentChunk,
                    Document.filename,
                    DocumentChunk.embedding.cosine_distance(query_emb).label("distance"),
                )
                .join(Document, Document.id == DocumentChunk.document_id)
                .where(Document.user_id == user_id, Document.status == DocumentStatus.READY)
                .order_by("distance")
                .limit(SEMANTIC_CANDIDATES)
            )
            result = await db.execute(stmt)
            for chunk, filename, distance in result.all():
                key = (chunk.document_id, chunk.chunk_index)
                similarity = 1.0 - float(distance or 1.0)
                parts[key] = {
                    "document_id": chunk.document_id,
                    "filename": filename,
                    "chunk_index": chunk.chunk_index,
                    "content": chunk.content,
                    "score": max(0.0, similarity),
                }

        # 2) Keyword (Postgres ILIKE OR-over-terms, scored by term hits)
        if mode in ("hybrid", "keyword"):
            terms = [t for t in _tokenize(query) if len(t) >= 3][:8]
            if terms:
                conditions = [DocumentChunk.content.ilike(f"%{t}%") for t in terms]
                from sqlalchemy import or_, func

                stmt = (
                    select(
                        DocumentChunk,
                        Document.filename,
                    )
                    .join(Document, Document.id == DocumentChunk.document_id)
                    .where(Document.user_id == user_id, Document.status == DocumentStatus.READY)
                    .where(or_(*conditions))
                    .limit(KEYWORD_CANDIDATES)
                )
                result = await db.execute(stmt)
                q_terms = set(terms)
                for chunk, filename in result.all():
                    key = (chunk.document_id, chunk.chunk_index)
                    lowered = chunk.content.lower()
                    hits = sum(1 for t in q_terms if t in lowered)
                    kw_score = hits / max(1, len(q_terms))
                    if key in parts:
                        # Blend with existing semantic score
                        parts[key]["score"] = (
                            HYBRID_ALPHA * parts[key]["score"] + (1 - HYBRID_ALPHA) * kw_score
                        )
                    else:
                        parts[key] = {
                            "document_id": chunk.document_id,
                            "filename": filename,
                            "chunk_index": chunk.chunk_index,
                            "content": chunk.content,
                            "score": kw_score,
                        }

        ranked = sorted(parts.values(), key=lambda p: p["score"], reverse=True)[:top_k]

        # Rerank blended hybrid results above the floor with keyword re-hits (cheap reranker)
        if mode == "hybrid" and ranked:
            q_terms = set(t for t in _tokenize(query) if len(t) >= 3)
            for item in ranked:
                lowered = item["content"].lower()
                re_hits = sum(1 for t in q_terms if t in lowered)
                if re_hits:
                    item["score"] = min(1.0, item["score"] + 0.05 * re_hits)
            ranked = sorted(ranked, key=lambda p: p["score"], reverse=True)

        return ranked

    # -- context / prompt building --------------------------------------------
    def build_context_block(self, results: list[dict]) -> str:
        """Format retrieved chunks into a numbered context block for the LLM."""
        if not results:
            return ""
        lines: list[str] = []
        used = 0
        for i, r in enumerate(results, 1):
            piece = f"[{i}] (source: {r['filename']}, chunk {r['chunk_index']})\n{r['content']}"
            if used + len(piece) > MAX_CONTEXT_CHARS:
                break
            lines.append(piece)
            used += len(piece)
        return "\n\n".join(lines)

    def build_rag_system_prompt(self, context_block: str) -> str:
        prompt = SYSTEM_PROMPT
        if context_block:
            prompt += (
                f"\n\nKNOWLEDGE BASE CONTEXT (cite as [source: filename]):\n{context_block}"
            )
        else:
            prompt += "\n\n(The knowledge base has no relevant documents for this question.)"
        return prompt

    async def build_chat_context(self, db: AsyncSession, *, user_id: uuid.UUID, query: str) -> dict:
        """Convenience wrapper used by the chat endpoint: retrieve + build prompt context."""
        results = await self.retrieve(db, user_id=user_id, query=query)
        context_block = self.build_context_block(results)
        sources = [
            {"document_id": str(r["document_id"]), "filename": r["filename"], "chunk_index": r["chunk_index"], "score": round(r["score"], 4)}
            for r in results
        ]
        return {"retrieved_context": context_block, "sources": sources}


rag_service = RagService()
